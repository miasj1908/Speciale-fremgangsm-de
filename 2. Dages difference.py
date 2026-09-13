import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
import shutil

INPUT_FILE  = r"C:\Users\miasj\PythonProjects\equity_reports_update.xlsx"
OUTPUT_FILE = r"C:\Users\miasj\PythonProjects\equity_reports_update_dagsdiff1.xlsx"

# --- Indlæs data ---
df_reports = pd.read_excel(INPUT_FILE, sheet_name="Orginal Reports")
df_cf      = pd.read_excel(INPUT_FILE, sheet_name="Original Cash Flows")

# Brug kun unikke (Company, Date) kombinationer fra CF-sheet
cf_unique = df_cf[["Company", "Date"]].drop_duplicates().copy()
cf_unique["cf_year"] = cf_unique["Date"].dt.year

df_reports["rep_year"] = df_reports["report_date"].dt.year

# --- Beregn nærmeste CF-dato inden for samme år ---
nearest_cf_dates = []
days_diffs       = []

for _, rep in df_reports.iterrows():
    candidates = cf_unique[
        (cf_unique["Company"] == rep["Company"]) &
        (cf_unique["cf_year"] == rep["rep_year"])
    ]

    if candidates.empty:
        nearest_cf_dates.append(pd.NaT)
        days_diffs.append(np.nan)
    else:
        diffs  = (candidates["Date"] - rep["report_date"]).abs()
        nearest = candidates.loc[diffs.idxmin(), "Date"]
        nearest_cf_dates.append(nearest)
        # Signed: positiv = CF efter equity report
        days_diffs.append((nearest - rep["report_date"]).days)

df_reports["nearest_cf_date"] = nearest_cf_dates
df_reports["days_diff"]        = days_diffs

# --- Gem til Excel med de nye kolonner ---
shutil.copy(INPUT_FILE, OUTPUT_FILE)

wb = load_workbook(OUTPUT_FILE)
ws = wb["Orginal Reports"]

# Find første ledige kolonne
max_col = ws.max_column + 1

# Header-styling
header_fill = PatternFill("solid", start_color="4472C4")
header_font = Font(bold=True, color="FFFFFF")

ws.cell(row=1, column=max_col,     value="nearest_cf_date").fill = header_fill
ws.cell(row=1, column=max_col,     value="nearest_cf_date").font = header_font
ws.cell(row=1, column=max_col + 1, value="days_diff").fill       = header_fill
ws.cell(row=1, column=max_col + 1, value="days_diff").font       = header_font

# Skriv værdier
for i, (cf_date, diff) in enumerate(zip(nearest_cf_dates, days_diffs), start=2):
    ws.cell(row=i, column=max_col,     value=cf_date if pd.notna(cf_date) else None)
    ws.cell(row=i, column=max_col + 1, value=int(diff) if pd.notna(diff) else None)

# Datoformat på nearest_cf_date kolonnen
from openpyxl.styles import numbers
col_letter = ws.cell(row=1, column=max_col).column_letter
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=max_col, max_col=max_col):
    for cell in row:
        cell.number_format = "DD-MM-YYYY"

wb.save(OUTPUT_FILE)

# --- Print opsummering ---
matched = df_reports["days_diff"].notna()
print(f"Equity reports i alt:              {len(df_reports)}")
print(f"Med CF-match samme år:             {matched.sum()}")
print(f"Uden CF-match samme år:            {(~matched).sum()}")
print()
print("Dagsdifference (signed, positiv = CF efter report):")
print(df_reports.loc[matched, "days_diff"].describe().round(1))
print()
print(f"Output gemt til: {OUTPUT_FILE}")

# --- Pivot-tabeller ---
df_short = pd.read_excel(INPUT_FILE, sheet_name="Short List Corrections")
df_short["exchange"] = df_short["Ticker"].str.split("_").str[0]
company_exchange = df_short[["Name", "exchange"]].drop_duplicates(subset="Name")
df_reports = df_reports.merge(company_exchange, left_on="Company", right_on="Name", how="left")

# Tabel 1: Sektor x År
tabel1 = df_reports.pivot_table(index="sector", columns="rep_year", aggfunc="size", fill_value=0)
tabel1["Total"] = tabel1.sum(axis=1)
tabel1.loc["Total"] = tabel1.sum()
print("\n=== Tabel 1: Equity reports per år og sektor ===")
print(tabel1.to_string())

# Tabel 2: Børsnotering x År
tabel2 = df_reports.pivot_table(index="exchange", columns="rep_year", aggfunc="size", fill_value=0)
tabel2["Total"] = tabel2.sum(axis=1)
tabel2.loc["Total"] = tabel2.sum()
print("\n=== Tabel 2: Equity reports per år og børsnotering ===")
print(tabel2.to_string())