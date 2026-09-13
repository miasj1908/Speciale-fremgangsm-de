"""
Scanner hele CBS mappen for at finde den nyeste laesbare cashflow model per selskab.
Bruger xlrd som primær metode og win32com (Excel) som backup for UTF-16 filer.

Krav:
    pip install xlrd psycopg2-binary pandas pywin32 openpyxl
"""

import pandas as pd
import xlrd
import openpyxl
import psycopg2
import os
import re
import tempfile
from datetime import datetime

LOKAL_ROD_STI = r"C:\Users\miasj\CBS - Copenhagen Business School\Caroline Thøisen Larsen - Data - Forecasting downloadet"

DB_HOST    = "localhost"
DB_PORT    = 5432
DB_BRUGER  = "postgres"
DB_KODEORD = "Miamia97!"
DB_NAVN    = "forecast_studie"

INCOME_STATEMENT = {
    "Revenue": "Revenue", "Cost of Goods Sold": "COGS",
    "Gross Profit": "Gross_Profit",
    "Selling, General, and Administrative Expenses": "SGA",
    "Other Operating Expense (Income)": "Other_Opex",
    "Depreciation & Amortization": "DA",
    "Operating Income (ex charges)": "EBIT",
    "Restructuring & Other Cash Charges": "Restructuring",
    "Impairment Charges": "Impairment",
    "Other Non-Cash (Income) / Charges": "Other_NonCash",
    "Operating Income (incl charges)": "EBIT_incl_charges",
    "Interest Expense": "Interest_Expense",
    "Interest Income": "Interest_Income",
    "Pre-Tax Income": "Pre_Tax_Income",
    "Income Tax Expense": "Tax_Expense",
    "Other After-Tax Cash Gains (Losses)": "Other_AfterTax_Cash",
    "Other After-Tax Non-Cash Gains (Losses)": "Other_AfterTax_NonCash",
    "(Minority Interest)": "Minority_Interest",
    "(Preferred Dividends)": "Preferred_Dividends",
    "Net Income": "Net_Income",
    "Weighted Average Diluted Shares Outstanding": "Shares_Diluted",
    "Diluted Earnings Per Share (GAAP)": "EPS_GAAP",
    "Adjustments to Net Income": "EPS_Adjustments",
    "Adjusted Net Income": "Net_Income_Adjusted",
    "Diluted Earnings Per Share (Adjusted)": "EPS_Adjusted",
    "Regular Dividends Per Share": "DPS_Regular",
    "Special Dividends Per Share": "DPS_Special",
    "(Total Common Dividends)": "Total_Dividends",
    "EBITDA": "EBITDA", "Adjusted EBITDA": "EBITDA_Adjusted",
}

BALANCE_SHEET = {
    "Cash and Equivalents": "Cash", "Investments": "Investments",
    "Accounts Receivable": "Accounts_Receivable", "Inventory": "Inventory",
    "Deferred Tax Assets (Current)": "DTA_Current",
    "Other Short-Term Assets": "Other_ST_Assets",
    "Current Assets": "Current_Assets",
    "Net Property, Plant, and Equipment": "Net_PPE",
    "Goodwill": "Goodwill", "Other Intangibles": "Other_Intangibles",
    "Deferred Tax Assets (Long-Term)": "DTA_LongTerm",
    "Other Long-Term Operating Assets": "Other_LT_Operating_Assets",
    "Long-Term Non-Operating Assets": "Other_LT_NonOp_Assets",
    "Total Assets": "Total_Assets", "Accounts Payable": "Accounts_Payable",
    "Short-Term Debt": "ST_Debt",
    "Deferred Tax Liabilities (Current)": "DTL_Current",
    "Other Short-Term Liabilities": "Other_ST_Liabilities",
    "Current Liabilities": "Current_Liabilities",
    "Long-Term Debt": "LT_Debt",
    "Deferred Tax Liabilities (Long-Term)": "DTL_LongTerm",
    "Other Long-Term Operating Liabilities": "Other_LT_Operating_Liab",
    "Long-Term Non-Operating Liabilities": "Other_LT_NonOp_Liab",
    "Total Liabilities": "Total_Liabilities",
    "Preferred Stock": "Preferred_Stock", "Common Stock": "Common_Stock",
    "Additional Paid-In Capital": "APIC",
    "Retained Earnings (Deficit)": "Retained_Earnings",
    "(Treasury Stock)": "Treasury_Stock", "Other Equity": "Other_Equity",
    "Shareholder's Equity": "Shareholders_Equity",
    "Minority Interest": "Minority_Interest_BS",
    "Total Equity": "Total_Equity",
}

CASH_FLOW = {
    "Net Income": "CF_Net_Income", "Depreciation": "Depreciation",
    "Amortization": "Amortization", "Stock-Based Compensation": "SBC",
    "Impairment of Goodwill": "Impairment_Goodwill",
    "Impairment of Other Intangibles": "Impairment_Intangibles",
    "Deferred Taxes": "Deferred_Taxes",
    "Other Non-Cash Adjustments": "Other_NonCash_CF",
    "(Increase) Decrease in Accounts Receivable": "Change_AR",
    "(Increase) Decrease in Inventory": "Change_Inventory",
    "Change in Other Short-Term Assets": "Change_Other_ST_Assets",
    "Increase (Decrease) in Accounts Payable": "Change_AP",
    "Change in Other Short-Term Liabilities": "Change_Other_ST_Liab",
    "Cash From Operations": "CFO", "(Capital Expenditures)": "Capex",
    "Net (Acquisitions), Asset Sales, and Disposals": "Net_Acquisitions",
    "Net (Purchases) Sales of Investments": "Net_Investments",
    "Other Investing Cash Flow": "Other_Investing_CF",
    "Cash From Investing": "CFI",
    "Common Stock Issuance or (Repurchase)": "Stock_Issuance_Repurchase",
    "Common Stock (Dividends)": "Dividends_Paid",
    "Short-Term Debt Issuance (Retirement)": "ST_Debt_Change",
    "Long-Term Debt Issuance (Retirement)": "LT_Debt_Change",
    "Other Financing Cash Flows": "Other_Financing_CF",
    "Cash From Financing": "CFF",
    "Exchange Rates, Discontinued Ops, etc. (net)": "FX_Other",
    "Net Change in Cash": "Net_Change_Cash",
}

SEKTIONER = {"income": INCOME_STATEMENT, "balance": BALANCE_SHEET, "cashflow": CASH_FLOW}
ALL_ITEMS = {**INCOME_STATEMENT, **BALANCE_SHEET, **CASH_FLOW}
DOBBELT_LABELS = {"Net Income": ["Net_Income", "CF_Net_Income"]}


def udtræk_ticker(filnavn):
    m = re.match(r'^(.+)_CashFlowModel_\d{8}\.xls$', filnavn)
    return m.group(1) if m else None


def scan_alle_filer(rod_sti):
    print(f"Scanner: {rod_sti}")
    alle_filer = []
    for dirpath, dirnames, filenames in os.walk(rod_sti):
        for fil in filenames:
            if fil.endswith('.xls') and 'CashFlowModel' in fil:
                ticker = udtræk_ticker(fil)
                dato_match = re.search(r'_(\d{8})\.xls$', fil)
                if ticker and dato_match:
                    try:
                        dato = datetime.strptime(dato_match.group(1), '%Y%m%d')
                        alle_filer.append({
                            'ticker': ticker,
                            'fil': fil,
                            'sti': os.path.join(dirpath, fil),
                            'dato': dato,
                        })
                    except Exception:
                        pass
    print(f"Fandt {len(alle_filer)} modeller")
    return alle_filer


def find_top_per_ticker(alle_filer, antal=50):
    df = pd.DataFrame(alle_filer)
    top = (
        df.sort_values('dato', ascending=False)
        .groupby('ticker')
        .head(antal)
        .reset_index(drop=True)
    )
    print(f"Unikke tickers: {df['ticker'].nunique()} (op til {antal} kandidater per ticker)")
    return top


def find_historiske_kolonner(ws_rows):
    """Finder historiske kolonner fra en liste af rækker."""
    offset_row = year_row = None
    for row in ws_rows[:10]:
        if -10 in row or -10.0 in row:
            offset_row = row
        if any(isinstance(v, (int, float)) and 2000 < v < 2035 for v in row):
            year_row = row
    if not offset_row or not year_row:
        return {}
    kol = {}
    for col_idx, (offset, year) in enumerate(zip(offset_row, year_row)):
        if isinstance(offset, (int, float)) and isinstance(year, (int, float)):
            if -10 <= int(offset) <= -1 and int(year) >= 2011:
                kol[int(year)] = col_idx
    return kol


def find_labels_fra_rows(alle_rows):
    """Scanner kolonne E (index 4) for labels."""
    forekomster = {}
    for row in alle_rows[:600]:
        if len(row) >= 5 and row[4] and isinstance(row[4], str):
            label = row[4].strip()
            if label in ALL_ITEMS:
                forekomster.setdefault(label, []).append(alle_rows.index(row))
    return forekomster


def udtræk_vaerdier_fra_rows(alle_rows, forekomster, kol_mapping):
    res = {}
    for excel_label, vores_navn in ALL_ITEMS.items():
        raekker = forekomster.get(excel_label, [])
        if excel_label in DOBBELT_LABELS:
            for i, navn in enumerate(DOBBELT_LABELS[excel_label]):
                row_idx = raekker[i] if i < len(raekker) else (raekker[0] if raekker else None)
                res[navn] = {}
                if row_idx is not None and row_idx < len(alle_rows):
                    row = alle_rows[row_idx]
                    for aar, col in kol_mapping.items():
                        v = row[col] if col < len(row) else None
                        res[navn][aar] = None if isinstance(v, str) else v
        else:
            res[vores_navn] = {}
            if raekker and raekker[0] < len(alle_rows):
                row = alle_rows[raekker[0]]
                for aar, col in kol_mapping.items():
                    v = row[col] if col < len(row) else None
                    res[vores_navn][aar] = None if isinstance(v, str) else v
    return res


def hent_rows_xlrd(filsti):
    """Indlaes med xlrd og returner alle rækker fra Inputs sheet."""
    try:
        wb = xlrd.open_workbook(filsti)
        for navn in ['Inputs', 'inputs', 'INPUTS']:
            if navn in wb.sheet_names():
                ws = wb.sheet_by_name(navn)
                return [ws.row_values(i) for i in range(min(600, ws.nrows))]
    except Exception:
        pass
    return None


def hent_rows_win32com(filsti):
    """Indlaes med Excel via win32com og returner alle rækker fra Inputs sheet."""
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        temp_dir = tempfile.mkdtemp()
        temp_xlsx = os.path.join(temp_dir, 'temp.xlsx')

        excel = win32com.client.Dispatch('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False

        wb = excel.Workbooks.Open(os.path.abspath(filsti))
        wb.SaveAs(temp_xlsx, FileFormat=51)
        wb.Close(False)
        excel.Quit()
        pythoncom.CoUninitialize()

        wb_px = openpyxl.load_workbook(temp_xlsx, read_only=True, data_only=True)
        for navn in ['Inputs', 'inputs', 'INPUTS']:
            if navn in wb_px.sheetnames:
                ws = wb_px[navn]
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 600:
                        break
                    rows.append([v if v is not None else '' for v in row])
                try:
                    os.remove(temp_xlsx)
                    os.rmdir(temp_dir)
                except Exception:
                    pass
                return rows

    except Exception:
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
    return None


def forsøg_indlaes(filsti, brug_win32=False):
    """Prøver xlrd først, derefter win32com hvis ønsket."""
    rows = hent_rows_xlrd(filsti)
    if rows is not None:
        return rows, 'xlrd'
    if brug_win32:
        rows = hent_rows_win32com(filsti)
        if rows is not None:
            return rows, 'win32com'
    return None, None


def opret_forbindelse():
    return psycopg2.connect(host=DB_HOST, port=DB_PORT,
                            user=DB_BRUGER, password=DB_KODEORD, dbname=DB_NAVN)


def gem_realiserede(cur, selskab_id, sektion, vaerdier):
    cur.execute("""
        SELECT forecast_id, første_prognose_år, sidste_prognose_år
        FROM forecasts WHERE selskab_id = %s
    """, (selskab_id,))
    forecasts = cur.fetchall()
    rows = []
    for fid, fp, sp in forecasts:
        for item, aar_dict in vaerdier.items():
            for aar, val in aar_dict.items():
                if val is not None and val != "" and fp <= aar <= sp:
                    try:
                        rows.append((fid, sektion, 'realiseret', item, aar, float(val)))
                    except Exception:
                        pass
    if rows:
        cur.executemany("""
            INSERT INTO regnskabsdata (forecast_id, sektion, datakilde, line_item, årstal, værdi)
            VALUES (%s, %s, %s, %s, %s, %s) ON CONFLICT DO NOTHING
        """, rows)
    return len(rows)


def koer():
    alle_filer = scan_alle_filer(LOKAL_ROD_STI)
    top = find_top_per_ticker(alle_filer, antal=50)

    conn = opret_forbindelse()
    cur = conn.cursor()

    cur.execute("SELECT ticker, selskab_id FROM selskaber")
    db_tickers = {row[0]: row[1] for row in cur.fetchall()}
    print(f"Tickers i database: {len(db_tickers)}")

    kandidater = top.groupby('ticker').apply(
        lambda x: x.sort_values('dato', ascending=False)[['sti', 'dato']].values.tolist()
    ).to_dict()

    succes = fejl_alle = ikke_i_db = total = 0
    win32_brugt = 0

    for ticker, fil_kandidater in kandidater.items():
        if ticker not in db_tickers:
            ikke_i_db += 1
            continue

        selskab_id = db_tickers[ticker]

        # Fase 1: Prøv alle kandidater med xlrd
        rows = None
        brugt_dato = None
        for filsti, dato in fil_kandidater:
            rows, metode = forsøg_indlaes(filsti, brug_win32=False)
            if rows is not None:
                brugt_dato = dato
                break

        # Fase 2: Hvis alle xlrd fejler, prøv win32com på de 3 nyeste
        if rows is None:
            print(f"  xlrd fejlede for {ticker} - prøver win32com...")
            for filsti, dato in fil_kandidater[:3]:
                rows, metode = forsøg_indlaes(filsti, brug_win32=True)
                if rows is not None:
                    brugt_dato = dato
                    win32_brugt += 1
                    print(f"  win32com lykkedes: {ticker} ({dato.strftime('%Y-%m-%d')})")
                    break

        if rows is None:
            print(f"  ALLE FEJL: {ticker}")
            fejl_alle += 1
            continue

        if succes % 20 == 0:
            print(f"  Match {succes}: {ticker} ({brugt_dato.strftime('%Y-%m-%d')})")
            conn.commit()

        try:
            kol = find_historiske_kolonner(rows)
            if not kol:
                fejl_alle += 1
                continue

            # Slet gamle realiserede tal
            cur.execute("""
                DELETE FROM regnskabsdata
                WHERE forecast_id IN (SELECT forecast_id FROM forecasts WHERE selskab_id = %s)
                AND datakilde = 'realiseret'
            """, (selskab_id,))

            forekomster = find_labels_fra_rows(rows)
            realiserede = udtræk_vaerdier_fra_rows(rows, forekomster, kol)

            for sek_navn, sek_items in SEKTIONER.items():
                sek = {k: v for k, v in realiserede.items() if k in sek_items.values()}
                total += gem_realiserede(cur, selskab_id, sek_navn, sek)

            succes += 1

        except Exception as e:
            conn.rollback()
            fejl_alle += 1

    conn.commit()

    # Evaluerbarhed
    cur.execute("""
        SELECT
            COUNT(*) FILTER (WHERE n >= 1),
            COUNT(*) FILTER (WHERE n >= 2),
            COUNT(*) FILTER (WHERE n >= 3),
            COUNT(*)
        FROM (
            SELECT f.forecast_id, COUNT(DISTINCT r.årstal) as n
            FROM forecasts f
            LEFT JOIN regnskabsdata r ON r.forecast_id = f.forecast_id
                AND r.datakilde = 'realiseret' AND r.line_item = 'EPS_Adjusted'
            GROUP BY f.forecast_id
        ) s
    """)
    t1, t2, t3, tot = cur.fetchone()

    print(f"\n=== EVALUERBARHED ===")
    print(f"t+1: {t1}/{tot} ({t1/tot*100:.1f}%)")
    print(f"t+2: {t2}/{tot} ({t2/tot*100:.1f}%)")
    print(f"t+3: {t3}/{tot} ({t3/tot*100:.1f}%)")

    cur.close()
    conn.close()

    print(f"\nSucces:          {succes}")
    print(f"Win32com brugt:  {win32_brugt}")
    print(f"Alle fejlede:    {fejl_alle}")
    print(f"Ikke i database: {ikke_i_db}")
    print(f"Raekker gemt:    {total}")


if __name__ == "__main__":
    start = datetime.now()
    koer()
    slut = datetime.now()
    print(f"Tid: {(slut-start).seconds//60}m {(slut-start).seconds%60}s")
